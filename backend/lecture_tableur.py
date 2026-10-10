"""Lecture des fichiers d'import catalogue : CSV, XLSX et XLS (ticket #86).

L'assistant d'import (page Catalogues → Importer) n'acceptait que le CSV. Ce
module lit aussi les classeurs Excel et renvoie le MEME objet qu'avant (un
DataFrame pandas dont toutes les cellules sont des chaînes), pour que le
mapping, les contrôles et l'activation restent inchangés.

Garde-fous repris du script d'import lourd
(scripts/fournisseur/import_catalogue_lourd.py), adaptés au plafond de
l'upload web (MAX_UPLOAD_SIZE, 15 Mo par défaut) :

- le type est lu dans les premiers octets, jamais dans l'extension seule ;
- XLSX : macros refusées (vbaProject.bin), nombre d'entrées de l'archive,
  taille décompressée et taux de compression bornés (zip bomb) ;
- feuilles, lignes, colonnes et longueur de cellule plafonnées ;
- valeurs calculées seulement (data_only) : une formule n'est jamais
  exécutée ni recopiée, seul son dernier résultat enregistré est lu ;
- l'en-tête est cherché dans les 30 premières lignes (titres, logos et lignes
  vides au-dessus du tableau sont ignorés).
"""
from __future__ import annotations

import datetime as _dt
import io
import zipfile

# Plafonds. Le fichier envoyé fait au plus MAX_UPLOAD_SIZE (15 Mo) ; un XLSX
# de 15 Mo se décompresse en général en 100 à 150 Mo.
XLSX_MAX_ENTREES = 5000
XLSX_MAX_DECOMPRESSE = 400 * 1024 * 1024
XLSX_RATIO_MAX = 150            # par entrée, au-delà de 20 Mo décompressés
MAX_FEUILLES = 50
MAX_LIGNES = 1_000_000
MAX_COLONNES = 500
MAX_CELLULE = 32_000
LIGNES_ENTETE_MAX = 30

SIGNATURE_XLSX = b"PK\x03\x04"
SIGNATURE_XLS = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"   # conteneur OLE (Excel 97-2003)


class FichierRefuse(ValueError):
    """Fichier refusé, avec un message lisible par l'utilisateur."""


def format_fichier(contenu: bytes, nom: str | None = None) -> str:
    """'xlsx', 'xls' ou 'csv', d'après le contenu (l'extension ne suffit pas)."""
    tete = contenu[:8]
    ext = (nom or "").lower().rsplit(".", 1)[-1] if nom and "." in nom else ""
    if tete.startswith(SIGNATURE_XLSX):
        return "xlsx"
    if tete.startswith(SIGNATURE_XLS):
        return "xls"
    if ext in ("xlsx", "xlsm", "xls"):
        raise FichierRefuse(f"Le fichier porte l'extension .{ext} mais son contenu n'est pas un classeur Excel.")
    if b"\x00" in contenu[:4096]:
        raise FichierRefuse("Ce fichier n'est ni un CSV texte ni un classeur Excel.")
    return "csv"


def controle_xlsx(contenu: bytes) -> dict:
    """Refuse macros, archives anormales et zip bombs. Renvoie quelques infos."""
    try:
        z = zipfile.ZipFile(io.BytesIO(contenu))
    except zipfile.BadZipFile as exc:
        raise FichierRefuse("Classeur Excel illisible (archive endommagée).") from exc
    with z:
        entrees = z.infolist()
        if len(entrees) > XLSX_MAX_ENTREES:
            raise FichierRefuse(f"Classeur suspect : {len(entrees)} éléments dans l'archive.")
        noms = {e.filename.lower() for e in entrees}
        if any(n.endswith("vbaproject.bin") for n in noms):
            raise FichierRefuse("Le classeur contient des macros : enregistrez-le en .xlsx sans macros, puis réessayez.")
        if "xl/workbook.xml" not in noms:
            raise FichierRefuse("Ce fichier n'est pas un classeur Excel (.xlsx).")
        total = 0
        for e in entrees:
            total += e.file_size
            if e.compress_size and e.file_size > 20 * 1024 * 1024 \
                    and e.file_size / e.compress_size > XLSX_RATIO_MAX:
                raise FichierRefuse("Taux de compression anormal dans le classeur : fichier refusé.")
        if total > XLSX_MAX_DECOMPRESSE:
            raise FichierRefuse(f"Classeur trop volumineux une fois décompressé ({total // (1024 * 1024)} Mo).")
        return {"decompresse_octets": total,
                "liens_externes": any("externallink" in n for n in noms)}


def texte_cellule(v) -> str:
    """Valeur de cellule → chaîne, sans artefact (8.0 → « 8 », date → AAAA-MM-JJ)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float):
        if v != v:                      # NaN
            return ""
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        return format(v, ".15g")        # 8.66 et non 8.6600000000000001 ; 0.1+0.2 → 0.3
    if isinstance(v, int):
        return str(v)
    if isinstance(v, _dt.datetime):
        return v.date().isoformat() if v.time() == _dt.time() else v.isoformat(sep=" ")
    if isinstance(v, _dt.date):
        return v.isoformat()
    s = str(v).strip()
    return s[:MAX_CELLULE]


def _est_entete(ligne: list[str]) -> bool:
    """Au moins 2 cellules remplies, dont une majorité de libellés (pas des nombres)."""
    remplies = [c for c in ligne if c]
    if len(remplies) < 2:
        return False
    textes = [c for c in remplies if not c.replace(",", ".").replace(" ", "").lstrip("-").replace(".", "", 1).isdigit()]
    return len(textes) >= max(2, (len(remplies) + 1) // 2)


def _tableau(lignes, nom_feuille: str):
    """Itérable de lignes brutes → (en-tête, données, n° de ligne de l'en-tête)."""
    entete, n_entete, donnees = None, 0, []
    for n, brute in enumerate(lignes, start=1):
        if len(donnees) >= MAX_LIGNES:
            raise FichierRefuse(f"Feuille « {nom_feuille} » : plus de {MAX_LIGNES:,} lignes. Utilisez l'import lourd.".replace(",", " "))
        ligne = [texte_cellule(v) for v in (brute or ())]
        if entete is None:
            if n > LIGNES_ENTETE_MAX:
                break
            if _est_entete(ligne):
                entete, n_entete = ligne, n
            continue
        if any(ligne):
            donnees.append(ligne)
    if entete is None:
        return None
    # Colonnes vides en fin d'en-tête retirées ; libellés manquants ou en double nommés.
    while entete and not entete[-1]:
        entete.pop()
    if len(entete) > MAX_COLONNES:
        raise FichierRefuse(f"Feuille « {nom_feuille} » : {len(entete)} colonnes (maximum {MAX_COLONNES}).")
    vus: dict[str, int] = {}
    noms = []
    for i, h in enumerate(entete):
        h = h or f"colonne_{i + 1}"
        if h in vus:
            vus[h] += 1
            h = f"{h} ({vus[h]})"
        else:
            vus[h] = 1
        noms.append(h)
    largeur = len(noms)
    donnees = [(l + [""] * largeur)[:largeur] for l in donnees]
    return noms, donnees, n_entete


def _feuilles_xlsx(contenu: bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    try:
        noms = wb.sheetnames
        if len(noms) > MAX_FEUILLES:
            raise FichierRefuse(f"Classeur de {len(noms)} feuilles (maximum {MAX_FEUILLES}).")
        for nom in noms:
            yield nom, wb[nom].iter_rows(values_only=True)
    finally:
        wb.close()


def _feuilles_xls(contenu: bytes):
    import xlrd

    try:
        wb = xlrd.open_workbook(file_contents=contenu, on_demand=True)
    except Exception as exc:  # xlrd.XLRDError, struct.error…
        raise FichierRefuse("Classeur Excel 97-2003 (.xls) illisible : enregistrez-le en .xlsx, puis réessayez.") from exc
    try:
        if wb.nsheets > MAX_FEUILLES:
            raise FichierRefuse(f"Classeur de {wb.nsheets} feuilles (maximum {MAX_FEUILLES}).")
        for i in range(wb.nsheets):
            ws = wb.sheet_by_index(i)

            def lignes(ws=ws):
                for r in range(ws.nrows):
                    vals = []
                    for c in range(ws.ncols):
                        cell = ws.cell(r, c)
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            vals.append(xlrd.xldate_as_datetime(cell.value, wb.datemode))
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                            vals.append(None)
                        else:
                            vals.append(cell.value)
                    yield vals
            yield ws.name, lignes()
            wb.unload_sheet(i)
    finally:
        wb.release_resources()


def lire_classeur(contenu: bytes, fmt: str, feuille: str | None = None):
    """Lit un XLSX/XLS. Renvoie (DataFrame de chaînes, infos).

    Sans `feuille`, la première feuille qui contient un tableau est retenue.
    `infos` : format, feuilles (toutes), feuille retenue, ligne d'en-tête.
    """
    import pandas as pd

    infos: dict = {"format": fmt}
    if fmt == "xlsx":
        infos.update(controle_xlsx(contenu))
        source = _feuilles_xlsx(contenu)
    else:
        source = _feuilles_xls(contenu)

    noms, retenu = [], None
    for nom, lignes in source:
        noms.append(nom)
        if retenu is not None:
            continue
        if feuille is not None and nom != feuille:
            continue
        t = _tableau(lignes, nom)
        if t is not None:
            retenu = (nom, t)
    infos["feuilles"] = noms
    if feuille is not None and feuille not in noms:
        raise FichierRefuse(f"La feuille « {feuille} » n'existe pas dans ce classeur.")
    if retenu is None:
        raise FichierRefuse("Aucun tableau trouvé : il faut une ligne d'en-tête (au moins deux colonnes nommées) "
                            f"dans les {LIGNES_ENTETE_MAX} premières lignes d'une feuille.")
    nom, (colonnes, donnees, n_entete) = retenu
    infos["feuille"] = nom
    infos["ligne_entete"] = n_entete
    df = pd.DataFrame(donnees, columns=colonnes, dtype=str)
    return df, infos
